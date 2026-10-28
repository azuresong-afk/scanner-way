"""Рендер отчётов: HTML (Jinja2) -> PDF (WeasyPrint), DOCX (python-docx), JSON.

Разрезы: по каждому модулю раздельно, сводный по всем модулям, тренд между
прогонами (open/fixed/new/regression + динамика по severity). Все форматы
строятся из одной модели (`reporting.model.Report`), в каждой находке —
рекомендация по устранению.

Графика (дашборд руководителя, распределение по опасности, топ хостов, модули,
каскад изменений и тренд по прогонам) строится в `reporting.charts`: одна
геометрия — inline-SVG в HTML/PDF (без JS и внешних ресурсов) и PNG в DOCX.

Безопасность: данные находок (evidence, баннеры сервисов) приходят с целей и
считаются недоверенными — HTML экранируется (autoescape), WeasyPrint не
загружает внешние ресурсы, из DOCX вырезаются недопустимые в XML символы.
"""
from __future__ import annotations

import json
import re
from collections.abc import Sequence
from dataclasses import dataclass
from enum import StrEnum
from io import BytesIO
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from markupsafe import Markup

from scanner.core.models import FindingStatus, Severity
from scanner.reporting import charts as ch
from scanner.reporting.charts import SEVERITY_HEX, STATUS_HEX, RunPoint, Scene
from scanner.reporting.diff import severity_counts
from scanner.reporting.model import (
    SEVERITY_LABELS,
    SEVERITY_ORDER,
    STATUS_LABELS,
    FindingView,
    Report,
)
from scanner.reporting.remediation import RefLink, describe_ref

TEMPLATES_DIR = Path(__file__).parent / "templates"

# Палитра DOCX — та же, что в HTML-шаблоне
DOCX_FONT = "Calibri"
INK, INK_2, MUTED, FAINT = "#101828", "#344054", "#667085", "#98a2b3"
LINE, PANEL, ACCENT, ACCENT_2 = "#eaecf0", "#f9fafb", "#1d2b4f", "#3e4c7a"
GOOD, GOOD_BG = "#067647", "#ecfdf3"

STATUS_ORDER: list[FindingStatus] = [
    FindingStatus.NEW, FindingStatus.REGRESSION, FindingStatus.OPEN, FindingStatus.FIXED,
]

SEVERITY_RGB: dict[Severity, tuple[int, int, int]] = {
    s: tuple(int(h[i:i + 2], 16) for i in (1, 3, 5))  # type: ignore[misc]
    for s, h in SEVERITY_HEX.items()
}


class ReportFormat(StrEnum):
    HTML = "html"
    PDF = "pdf"
    DOCX = "docx"
    JSON = "json"


CONTENT_TYPES: dict[ReportFormat, str] = {
    ReportFormat.HTML: "text/html; charset=utf-8",
    ReportFormat.PDF: "application/pdf",
    ReportFormat.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ReportFormat.JSON: "application/json",
}


def _env() -> Environment:
    env = Environment(
        loader=FileSystemLoader(TEMPLATES_DIR),
        autoescape=select_autoescape(["html", "j2"]),
        trim_blocks=True,
        lstrip_blocks=True,
    )
    env.globals.update(
        SEVERITY_ORDER=SEVERITY_ORDER, SEVERITY_LABELS=SEVERITY_LABELS,
        STATUS_ORDER=STATUS_ORDER, STATUS_LABELS=STATUS_LABELS,
        SEVERITY_HEX={str(k): v for k, v in SEVERITY_HEX.items()},
    )
    env.filters["links"] = _links
    return env


def _links(urls: list[str], limit: int | None = None) -> list[RefLink]:
    """Ссылки с подписями (бюллетень производителя, ОС, БДУ, NVD, CWE…)."""
    out = [link for u in urls if (link := describe_ref(u)) is not None]
    return out[:limit] if limit else out


def _rgb(hex_color: str):
    from docx.shared import RGBColor
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _deny_fetch(url: str, *args, **kwargs):
    raise ValueError(f"загрузка внешних ресурсов в отчёт запрещена: {url}")


# Символы, недопустимые в XML 1.0 (python-docx на них падает)
_XML_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def _x(text: object) -> str:
    return _XML_BAD.sub("�", "" if text is None else str(text))


@dataclass(slots=True)
class HostBlock:
    """Все находки одного хоста (раздел «Результаты по хостам»)."""
    host: str
    findings: list[FindingView]
    counts: dict[Severity, int]

    @property
    def total(self) -> int:
        return len(self.findings)

    @property
    def risk(self) -> Severity | None:
        return next((s for s in SEVERITY_ORDER if self.counts[s]), None)

    @property
    def ports(self) -> list[int]:
        return sorted({v.finding.port for v in self.findings if v.finding.port is not None})


def host_blocks(report: Report) -> list[HostBlock]:
    """Находки по хостам; хосты — по убыванию риска (как топ хостов в резюме)."""
    per: dict[str, list[FindingView]] = {}
    for v in report.combined.findings:      # уже отсортированы по опасности
        per.setdefault(v.finding.host, []).append(v)
    blocks = [HostBlock(h, vs, severity_counts([v.finding for v in vs]))
              for h, vs in per.items()]
    blocks.sort(key=lambda b: tuple(-b.counts[s] for s in SEVERITY_ORDER) + (b.host,))
    return blocks


def _run_label(run_id: object, fallback: str) -> str:
    return fallback if run_id is None else f"Прогон {run_id}"


def default_history(report: Report) -> list[RunPoint]:
    """Две точки тренда (предыдущий -> текущий прогон) из самой модели отчёта."""
    t = report.combined.trend
    if t is None:
        return []
    c = report.context
    return [
        RunPoint(_run_label(c.previous_run_id, "Предыдущий"),
                 {d.severity: d.previous for d in t.severity}),
        RunPoint(_run_label(c.run_id, "Текущий"), {d.severity: d.current for d in t.severity}),
    ]


def build_charts(report: Report, history: Sequence[RunPoint] | None = None
                 ) -> dict[str, Scene]:
    """Диаграммы отчёта по ключам; набор зависит от данных (тренд — со 2-го прогона).

    history — точки по прогонам в хронологическом порядке (последняя — текущий
    прогон); если не задана, берутся предыдущий и текущий прогон из отчёта.
    """
    s = report.summary
    out: dict[str, Scene] = {"severity": ch.severity_donut(s.severity_counts)}
    if s.top_hosts:
        out["hosts"] = ch.stacked_bars(
            "hosts", "Наиболее уязвимые хосты по уровню опасности",
            [(h.host, h.counts) for h in s.top_hosts], width=340, label_w=110)
    if report.modules:
        out["modules"] = ch.stacked_bars(
            "modules", "Находки по модулям и уровню опасности",
            [(m.label, m.severity_counts) for m in report.modules], width=680, label_w=170)
    t = report.combined.trend
    if t is not None:
        k = t.counts
        out["waterfall"] = ch.waterfall(
            k[FindingStatus.OPEN] + k[FindingStatus.FIXED], k[FindingStatus.FIXED],
            k[FindingStatus.NEW], k[FindingStatus.REGRESSION])
        points = list(history) if history else default_history(report)
        out["trend"] = ch.trend_columns(points)
    return out


class ReportRenderer:
    def render(self, report: Report, fmt: ReportFormat | str,
               history: Sequence[RunPoint] | None = None) -> bytes:
        fmt = ReportFormat(fmt)
        if fmt is ReportFormat.HTML:
            return self.render_html(report, history).encode("utf-8")
        if fmt is ReportFormat.PDF:
            return self.render_pdf(report, history)
        if fmt is ReportFormat.DOCX:
            return self.render_docx(report, history)
        return self.render_json(report)

    # --- JSON ---------------------------------------------------------------
    def render_json(self, report: Report) -> bytes:
        return json.dumps(report.to_dict(), ensure_ascii=False, indent=2).encode("utf-8")

    # --- HTML / PDF ---------------------------------------------------------
    def render_html(self, report: Report, history: Sequence[RunPoint] | None = None) -> str:
        # SVG собран из экранированных строк (charts.py) — помечаем как безопасный
        charts = {k: Markup(v.to_svg()) for k, v in build_charts(report, history).items()}
        return _env().get_template("report.html.j2").render(
            r=report, charts=charts, hosts=host_blocks(report))

    def render_pdf(self, report: Report, history: Sequence[RunPoint] | None = None) -> bytes:
        try:
            from weasyprint import HTML  # тяжёлая зависимость (pango) — лениво
        except (ImportError, OSError) as e:
            raise RuntimeError(
                "PDF недоступен: нужен WeasyPrint и системные библиотеки Pango") from e
        return HTML(string=self.render_html(report, history),
                    url_fetcher=_deny_fetch).write_pdf()

    # --- DOCX ---------------------------------------------------------------
    def render_docx(self, report: Report, history: Sequence[RunPoint] | None = None) -> bytes:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_BREAK
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor

        doc = Document()
        for sec in doc.sections:
            sec.left_margin = sec.right_margin = Cm(2.0)
            sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
        style = doc.styles["Normal"]
        style.font.name = DOCX_FONT
        style.font.size = Pt(10)
        style.font.color.rgb = _rgb(INK_2)
        style.paragraph_format.space_after = Pt(4)
        for name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 12.5),
                           ("Heading 3", 11)):
            st = doc.styles[name]
            st.font.name = DOCX_FONT
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = _rgb(INK)
            st.element.rPr.rFonts.set(qn("w:asciiTheme"), DOCX_FONT)  # перебить тему
            st.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
            st.paragraph_format.space_after = Pt(6)

        page_cm = 17.0  # ширина текста A4 при полях 2 см
        scenes = build_charts(report, history) if ch.png_available() else {}

        def shade(el, fill: str):
            pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fill.lstrip("#"))
            pr.append(shd)

        def hairlines(t):
            """Только горизонтальные тонкие линии — как в HTML-отчёте."""
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                b = OxmlElement(f"w:{edge}")
                if edge in ("left", "right", "insideV"):
                    b.set(qn("w:val"), "nil")
                else:
                    b.set(qn("w:val"), "single")
                    b.set(qn("w:sz"), "4")
                    b.set(qn("w:color"), LINE.lstrip("#"))
                borders.append(b)
            t._tbl.tblPr.append(borders)

        def spacer(pt: float = 4):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.add_run().font.size = Pt(pt)

        def hyperlink(par, url: str, size: float = 8):
            """Кликабельная ссылка (python-docx не умеет их «из коробки»)."""
            rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
            h = OxmlElement("w:hyperlink")
            h.set(qn("r:id"), rid)
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            for tag, val in (("w:color", ACCENT_2.lstrip("#")), ("w:u", "single"),
                             ("w:sz", str(int(size * 2)))):
                el = OxmlElement(tag)
                el.set(qn("w:val"), val)
                rpr.append(el)
            r.append(rpr)
            t = OxmlElement("w:t")
            t.text = _x(url)
            r.append(t)
            h.append(r)
            par._p.append(h)

        def small(par, text: str, color: str = MUTED, size: float = 8.5, bold=False,
                  italic=False):
            run = par.add_run(_x(text))
            run.font.size = Pt(size)
            run.font.color.rgb = _rgb(color)
            run.bold = bold
            run.italic = italic
            return run

        def sev_cell(cell, sv: Severity):
            cell.text = ""
            r = cell.paragraphs[0].add_run("● " + SEVERITY_LABELS[sv])
            r.bold = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(*SEVERITY_RGB[sv])

        def table(header: list[str], rows: list[list[object]],
                  widths: list[float] | None = None, first_cm: float | None = None,
                  num_from: int | None = None):
            t = doc.add_table(rows=1, cols=len(header))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            hairlines(t)
            if widths is None and first_cm is not None:
                rest = (page_cm - first_cm) / (len(header) - 1)
                widths = [first_cm] + [rest] * (len(header) - 1)
            for cell, h in zip(t.rows[0].cells, header):
                cell.text = ""
                shade(cell._tc, PANEL)
                small(cell.paragraphs[0], h.upper(), MUTED, 7.5, bold=True)
            for row in rows:
                cells = t.add_row().cells
                for j, (cell, v) in enumerate(zip(cells, row)):
                    if isinstance(v, Severity):
                        sev_cell(cell, v)
                        continue
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(_x(v))
                    run.font.size = Pt(9)
                    if num_from is not None and j >= num_from:
                        cell.paragraphs[0].alignment = 2  # right
                        if v in (0, "0"):
                            run.font.color.rgb = _rgb(FAINT)
            if num_from is not None:
                for cell in t.rows[0].cells[num_from:]:
                    cell.paragraphs[0].alignment = 2
            if widths:
                t.autofit = False
                for col, w in zip(t.columns, widths):
                    col.width = Cm(w)  # tblGrid — его читает LibreOffice
                for row in t.rows:
                    for cell, w in zip(row.cells, widths):
                        cell.width = Cm(w)
            spacer()
            return t

        def picture(key: str, width_cm: float):
            if key in scenes:
                doc.add_picture(BytesIO(scenes[key].to_png()), width=Cm(width_cm))
                doc.paragraphs[-1].alignment = 1  # по центру

        def heading(num: int, text: str, page_break: bool = True):
            if page_break:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h = doc.add_heading(level=1)
            r = h.add_run(f"{num:02d}  ")
            r.font.color.rgb = _rgb(ACCENT_2)
            h.add_run(_x(text))

        def kpis(items: list[tuple[str, str, str | None]]):
            t = doc.add_table(rows=1, cols=len(items))
            hairlines(t)
            for cell, (label, value, color) in zip(t.rows[0].cells, items):
                cell.text = ""
                shade(cell._tc, PANEL)
                small(cell.paragraphs[0], label.upper(), MUTED, 7, bold=True)
                vp = cell.add_paragraph()
                vr = vp.add_run(_x(value))
                vr.bold = True
                vr.font.size = Pt(17)
                vr.font.color.rgb = _rgb(color or INK)
            spacer()

        c = report.context
        s = report.summary
        t = report.combined.trend

        # Обложка
        spacer(60)
        small(doc.add_paragraph(), "АНАЛИЗ ЗАЩИЩЁННОСТИ · РЕЗУЛЬТАТЫ СКАНИРОВАНИЯ",
              ACCENT_2, 8, bold=True)
        title = doc.add_paragraph(style="Title")
        title.add_run(_x(c.title)).font.color.rgb = _rgb(ACCENT)
        if c.customer:
            cp = doc.add_paragraph()
            cr = cp.add_run(_x(c.customer))
            cr.bold = True
            cr.font.size = Pt(14)
            cr.font.color.rgb = _rgb(INK)
        meta = []
        if c.engagement:
            meta.append(f"Договор / проект: {c.engagement}")
        if c.run_id is not None:
            meta.append(f"Прогон: {c.run_id}")
        if c.previous_run_id is not None:
            meta.append(f"Сравнение с прогоном: {c.previous_run_id}")
        meta.append(f"Сформирован: {c.generated_at:%d.%m.%Y %H:%M} UTC")
        small(doc.add_paragraph(), " · ".join(meta), MUTED, 9)
        spacer(24)
        kpis([("Уровень риска", s.risk_label,
               SEVERITY_HEX[s.risk_level] if s.risk_level else GOOD),
              ("Хостов с находками", str(s.hosts), None),
              ("Всего находок", str(s.total), None),
              ("Критических и высоких",
               str(s.severity_counts[Severity.CRITICAL] + s.severity_counts[Severity.HIGH]),
               None)])
        spacer(80)
        small(doc.add_paragraph(),
              "Конфиденциально. Отчёт сформирован по результатам авторизованного "
              "тестирования в рамках договора и предназначен только для заказчика.",
              FAINT, 8)

        # 1. Резюме
        heading(1, "Резюме для руководства")
        kp = [("Общий уровень риска", s.risk_label,
               SEVERITY_HEX[s.risk_level] if s.risk_level else GOOD),
              ("Всего находок", str(s.total), None),
              ("Хостов с находками", str(s.hosts), None)]
        if t:
            fresh = f"{t.counts[FindingStatus.NEW]} / {t.counts[FindingStatus.REGRESSION]}"
            kp.append(("Новые / повторные", fresh, None))
        kpis(kp)
        picture("severity", 10.0)
        if s.top_hosts:
            picture("hosts", 11.0)
        doc.add_heading("Ключевые выводы", level=2)
        for point in s.key_points:
            doc.add_paragraph(_x(point), style="List Bullet")
        doc.add_heading("Находки по модулям", level=2)
        picture("modules", page_cm)
        table(["Модуль"] + [SEVERITY_LABELS[x] for x in SEVERITY_ORDER] + ["Всего"],
              [[m.label] + [m.severity_counts[x] for x in SEVERITY_ORDER]
               + [len(m.findings)] for m in report.modules], first_cm=4.0, num_from=1)

        # 2. План устранения
        heading(2, "План устранения")
        if report.plan:
            small(doc.add_paragraph(), "Одинаковые меры объединены: одна строка — одно "
                  "действие и все хосты, где его нужно выполнить. Порядок — по уровню "
                  "опасности и числу затронутых хостов.", MUTED, 9)
            pt = table(["№", "Опасность", "Что сделать", "Где"],
                       [[i, p.severity, p.action, ", ".join(p.targets)]
                        for i, p in enumerate(report.plan, 1)],
                       widths=[0.8, 2.8, 9.6, 3.8])
            for row, p in zip(pt.rows[1:], report.plan):
                cell = row.cells[2]
                cell.paragraphs[0].runs[0].bold = True
                small(cell.add_paragraph(),
                      f"Источник: {p.source_label} · {'; '.join(p.titles)}", MUTED, 8)
                for link in _links(p.references, 4):
                    lp = cell.add_paragraph()
                    small(lp, f"{link.label}: ", INK_2, 8)
                    hyperlink(lp, link.url)
        else:
            doc.add_paragraph("Мер устранения не требуется.")

        # 3. Сравнение
        heading(3, "Сравнение с предыдущим прогоном")
        if t:
            vp = doc.add_paragraph()
            vr = vp.add_run(_x(t.verdict))
            vr.bold = True
            vr.font.color.rgb = _rgb(INK)
            kpis([(STATUS_LABELS[x], str(t.counts[x]), STATUS_HEX[x]) for x in STATUS_ORDER])
            picture("waterfall", 10.5)
            picture("trend", 10.0)
            table(["Уровень опасности", "Предыдущий", "Текущий", "Изменение"],
                  [[d.severity, d.previous, d.current, f"{d.delta:+d}" if d.delta else "0"]
                   for d in t.severity], first_cm=5.0, num_from=1)
            doc.add_heading("Динамика по модулям", level=2)
            table(["Модуль"] + [STATUS_LABELS[x] for x in STATUS_ORDER],
                  [[m.label] + [m.trend.counts[x] for x in STATUS_ORDER]
                   for m in report.modules if m.trend], first_cm=4.5, num_from=1)
            if t.fixed:
                doc.add_heading("Устранено с предыдущего прогона", level=2)
                table(["Опасность", "Находка", "Хост"],
                      [[v.finding.severity, v.finding.title, v.target] for v in t.fixed],
                      widths=[3.0, 10.0, 4.0])
        else:
            doc.add_paragraph("Предыдущий прогон не задан — это первый прогон, "
                              "сравнение будет доступно со следующего.")

        # 4. По хостам
        heading(4, "Результаты по хостам")
        hosts = host_blocks(report)
        if not hosts:
            doc.add_paragraph("Находок нет.")
        else:
            small(doc.add_paragraph(), "Все находки по каждому хосту. Хосты упорядочены "
                  "по уровню риска.", MUTED, 9)
        for b in hosts:
            h = doc.add_heading(level=2)
            h.add_run(_x(b.host))
            if b.risk:
                rr = h.add_run(f"   ● {SEVERITY_LABELS[b.risk]}")
                rr.font.size = Pt(9)
                rr.font.color.rgb = RGBColor(*SEVERITY_RGB[b.risk])
            facts = f"Находок: {b.total}"
            if b.ports:
                facts += " · Порты: " + ", ".join(map(str, b.ports))
            facts += " · " + ", ".join(f"{SEVERITY_LABELS[x].lower()}: {b.counts[x]}"
                                       for x in SEVERITY_ORDER if b.counts[x])
            small(doc.add_paragraph(), facts, MUTED, 8.5)
            header = ["Опасность", "Находка", "Порт", "Идентификаторы", "CVSS"]
            widths = [3.1, 6.6, 1.2, 4.6, 1.5]
            if t:
                header.append("Статус")
                widths = [3.1, 5.0, 1.1, 3.8, 1.2, 2.8]
            rows = []
            for v in b.findings:
                f = v.finding
                ids = "\n".join(list(f.cve_ids) + list(f.bdu_ids)) or (f.vuln_id or "—")
                row = [f.severity, f.title, f.port if f.port is not None else "—", ids,
                       f"{f.cvss:.1f}" if f.cvss is not None else "—"]
                if t:
                    row.append(v.status_label or "")
                rows.append(row)
            table(header, rows, widths=widths)

        # 5. Детально
        heading(5, "Детальные результаты для инженеров")
        for i, m in enumerate(report.modules, 1):
            doc.add_heading(_x(f"5.{i}. {m.label}"), level=2)
            if not m.findings:
                small(doc.add_paragraph(), "Открытых находок нет.", MUTED, 9)
            for _, _, group in m.by_severity:
                for v in group:
                    f = v.finding
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.keep_with_next = True
                    sr = p.add_run(f"● {SEVERITY_LABELS[f.severity]}   ")
                    sr.bold = True
                    sr.font.size = Pt(8.5)
                    sr.font.color.rgb = RGBColor(*SEVERITY_RGB[f.severity])
                    tr = p.add_run(_x(f.title))
                    tr.bold = True
                    tr.font.color.rgb = _rgb(INK)
                    if v.status_label:
                        small(p, f"  · {v.status_label}", MUTED, 8.5, italic=True)
                    details = [f"Хост: {v.target}"]
                    if f.vuln_id:
                        details.append(f"ID: {f.vuln_id}")
                    if f.cvss is not None:
                        details.append(f"CVSS: {f.cvss:.1f}")
                    if f.cve_ids:
                        details.append("CVE: " + ", ".join(f.cve_ids))
                    if f.bdu_ids:
                        details.append("БДУ: " + ", ".join(f.bdu_ids))
                    dp = doc.add_paragraph()
                    dp.paragraph_format.keep_with_next = True
                    small(dp, " · ".join(details), MUTED, 8.5)
                    if f.evidence:
                        ev = doc.add_paragraph()
                        shade(ev._p, PANEL)
                        r = ev.add_run(_x(f.evidence))
                        r.font.name = "Consolas"
                        r.font.size = Pt(8)
                    rem = v.remediation
                    rp = doc.add_paragraph()
                    shade(rp._p, GOOD_BG)
                    small(rp, "Рекомендация по устранению\n", GOOD, 8, bold=True)
                    rp.add_run(_x(rem.text))
                    note = f"Источник: {rem.source_label}."
                    if rem.fix_method:
                        note += f" Способ устранения: {rem.fix_method}."
                    small(rp, "\n" + note, MUTED, 8)
                    if rem.references:
                        small(rp, "\nСсылки на устранение:", GOOD, 8, bold=True)
                        for link in _links(rem.references):
                            small(rp, f"\n{link.label}: ", INK_2, 8, bold=True)
                            hyperlink(rp, link.url)

        spacer(10)
        small(doc.add_paragraph(), "Отчёт сформирован по результатам авторизованного "
              "тестирования в рамках договора. Порядок и сроки устранения определяются "
              "политикой заказчика.", MUTED, 8, italic=True)
        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()
